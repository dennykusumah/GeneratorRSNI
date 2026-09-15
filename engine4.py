"""Engine 4 — sisipkan halaman Prakata setelah Daftar isi.

Engine menerima keluaran Engine 3 dan mempertahankan susunan section:
1 Cover, 2 Copyright, 3 Daftar isi + Prakata + Introduction,
4 Content sampai Bibliography.
"""

from __future__ import annotations

from pipeline_utils import validate_docx, atomic_save_docx

import os
import re
from datetime import datetime
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


class PrakataPendahuluanEngine:
    FONT_NAME = "Arial"

    @staticmethod
    def _normalize(text: str) -> str:
        return re.sub(r"\s+", " ", text or "").strip()

    @classmethod
    def _key(cls, text: str) -> str:
        return cls._normalize(text).casefold()

    @staticmethod
    def _set_fonts(run, size: float = 11, bold=False, italic=False, color=None):
        run.font.name = "Arial"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), "Arial")

    @staticmethod
    def _section_before(paragraph) -> int:
        breaks = 0
        for sibling in paragraph._p.itersiblings(preceding=True):
            if sibling.tag != qn("w:p"):
                continue
            ppr = sibling.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                breaks += 1
        return breaks + 1

    def _ensure_judul_style(self, doc: Document):
        try:
            style = doc.styles["@Judul"]
        except KeyError:
            style = doc.styles.add_style("@Judul", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Arial"
        style.font.size = Pt(12)
        style.font.bold = True
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), "Arial")
        fmt = style.paragraph_format
        fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt.space_before = Pt(0)
        fmt.space_after = Pt(0)
        fmt.line_spacing = 1.0
        return style

    def _new_paragraph(self, doc, *, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        return p

    def _add_run(self, paragraph, text, *, italic=False, bold=False, size=11):
        run = paragraph.add_run(text)
        self._set_fonts(run, size=size, bold=bold, italic=italic)
        return run

    def _blank(self, doc):
        return self._new_paragraph(doc)

    def _page_break(self, doc):
        p = self._new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.LEFT)
        p.add_run().add_break(WD_BREAK.PAGE)
        return p

    def _bullet(self, doc, text):
        p = self._new_paragraph(doc)
        p.paragraph_format.left_indent = Cm(0.635)
        p.paragraph_format.first_line_indent = Cm(-0.635)
        self._add_run(p, "—    ")
        self._add_run(p, text)
        return p

    def _detect_content_title(self, doc: Document) -> str:
        scope_index = None
        for i, p in enumerate(doc.paragraphs):
            if re.match(r"^\s*1\s{1,4}Scope\b", self._normalize(p.text), re.I):
                scope_index = i
                break
        if scope_index is None:
            for p in doc.paragraphs:
                text = self._normalize(p.text)
                if (
                    text
                    and p.style.name.casefold() == "@judul".casefold()
                    and text.casefold() not in {
                        "introduction", "prakata", "daftar isi",
                        "bibliography", "bibliographie",
                    }
                ):
                    return text
            return "Judul Standar"
        for p in reversed(doc.paragraphs[:scope_index]):
            text = self._normalize(p.text)
            if text and text.casefold() not in {"introduction", "prakata", "daftar isi"}:
                return text
        return "Judul Standar"

    def _build_prakata(
        self, doc, sni_number, title_id, title_en, ref_standard, bsn_year,
        leading_page_break=True, trailing_page_break=True,
    ):
        items = []
        # Prakata wajib dimulai pada halaman baru setelah Daftar isi.
        if leading_page_break:
            items.append(self._page_break(doc))
        title = self._new_paragraph(doc, align=WD_ALIGN_PARAGRAPH.CENTER)
        title.style = self._ensure_judul_style(doc)
        title.add_run("Prakata")
        items.append(title)
        items.extend(self._blank(doc) for _ in range(3))

        p = self._new_paragraph(doc)
        self._add_run(p, f"{sni_number}, ")
        self._add_run(p, title_id, italic=True)
        self._add_run(
            p,
            ", merupakan standar yang disusun dengan jalur adopsi tingkat "
            "keselarasan identik dari ",
        )
        self._add_run(p, ref_standard)
        self._add_run(p, ", ")
        self._add_run(p, title_en, italic=True)
        self._add_run(
            p,
            f", dengan metode adopsi terjemahan dua bahasa dan ditetapkan "
            f"oleh BSN Tahun {bsn_year}.",
        )
        items.extend((p, self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(p, "Dalam Standar ini istilah “")
        self._add_run(p, "this International Standard", italic=True)
        self._add_run(p, f"” pada standar {ref_standard} yang diadopsi diganti dengan “")
        self._add_run(p, "this Standard", italic=True)
        self._add_run(p, "” dan diterjemahkan menjadi “Standar ini”.")
        items.extend((p, self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(
            p,
            "Terdapat standar yang dijadikan sebagai acuan normatif dalam "
            "Standar ini telah diadopsi menjadi SNI, yaitu:",
        )
        items.extend((p, self._blank(doc)))
        bullet_text = (
            "ISO/IEC XXXX-X:YYYY, ZZZZ, telah diadopsi dengan tingkat "
            "keselarasan identik menjadi SNI ISO/IEC XXXX-X:YYYY, ZZZZ"
        )
        items.extend((self._bullet(doc, bullet_text), self._blank(doc)))
        items.extend((self._bullet(doc, bullet_text), self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(
            p,
            "Standar ini disusun oleh Komite Teknis XX-YY, ZZZZ. Standar ini "
            "telah dibahas melalui rapat teknis dan disepakati dalam rapat "
            "konsensus pada tanggal XXXX di YYYY, yang dihadiri oleh para "
            "pemangku kepentingan (",
        )
        self._add_run(p, "stakeholders", italic=True)
        self._add_run(
            p,
            ") terkait yaitu perwakilan dari pemerintah, pelaku usaha, konsumen, "
            "dan pakar. Standar ini telah melalui tahap jajak pendapat pada "
            "tanggal XXXX sampai dengan YYYY dengan hasil akhir disetujui menjadi SNI.",
        )
        items.extend((p, self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(
            p,
            "Untuk menghindari kesalahan dalam penggunaan Standar ini, disarankan "
            "bagi pengguna standar menggunakan dokumen SNI yang dicetak dengan "
            "tinta berwarna (dapat mencantumkan kode tingkat warna ",
        )
        self._add_run(p, "Red Green Blue", italic=True)
        self._add_run(
            p,
            " (RGB) jika diperlukan untuk cetak gambar dengan warna yang lebih akurat).",
        )
        items.extend((p, self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(
            p,
            f"Apabila pengguna menemukan keraguan dalam Standar ini, maka "
            f"disarankan untuk melihat standar aslinya, yaitu {ref_standard}, "
            "dan/atau dokumen terkait lain yang menyertainya.",
        )
        items.extend((p, self._blank(doc)))

        p = self._new_paragraph(doc)
        self._add_run(
            p,
            "Perlu diperhatikan bahwa kemungkinan beberapa unsur dari Standar ini "
            "dapat berupa kekayaan intelektual. Namun selama proses perumusan SNI, "
            "Badan Standardisasi Nasional telah memperhatikan penyelesaian terhadap "
            "kemungkinan adanya kekayaan intelektual terkait substansi SNI. Apabila "
            "setelah penetapan SNI masih terdapat permasalahan terkait kekayaan "
            "intelektual, Badan Standardisasi Nasional tidak bertanggung jawab "
            "mengenai bukti, validitas, dan ruang lingkup dari kekayaan intelektual tersebut.",
        )
        items.append(p)
        if trailing_page_break:
            items.append(self._page_break(doc))
        return items

    def insert_prakata(
        self,
        input_docx: str,
        output_docx: str,
        sni_number: str = "SNI ISO/IEC XXXX:20XX",
        title_id: str = "",
        title_en: str = "",
        ref_standard: str = "",
        bsn_year: str = "",
        trailing_page_break: bool = True,
    ):
        try:
            if not input_docx or not os.path.isfile(input_docx):
                raise FileNotFoundError(f"File input tidak ditemukan: {input_docx}")
            doc = Document(input_docx)
            input_section_count = len(doc.sections)
            if input_section_count < 4:
                raise ValueError(
                    f"Output Engine 3 harus memiliki minimal 4 section; "
                    f"ditemukan {input_section_count}."
                )

            intro = next((p for p in doc.paragraphs if self._key(p.text) == "introduction"), None)
            toc = next((p for p in doc.paragraphs if self._key(p.text) == "daftar isi"), None)
            if toc is None:
                raise ValueError("Daftar isi tidak ditemukan.")
            if self._section_before(toc) != 3:
                raise ValueError("Daftar isi harus berada di section 3.")

            insertion_element = intro._p if intro is not None else None
            if intro is not None:
                if self._section_before(intro) != 3:
                    raise ValueError("Introduction harus berada di section 3.")
            else:
                boundaries = []
                for paragraph in doc.paragraphs:
                    ppr = paragraph._p.find(qn("w:pPr"))
                    if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                        boundaries.append(paragraph._p)
                if len(boundaries) < 3:
                    raise ValueError("Batas section front matter dan Content tidak ditemukan.")
                insertion_element = boundaries[2]

            # Idempoten jika dokumen yang sama diproses ulang.
            if any(self._key(p.text) == "prakata" for p in doc.paragraphs):
                atomic_save_docx(doc, output_docx)
                return True, output_docx

            detected_title = self._detect_content_title(doc)
            title_id = self._normalize(title_id or detected_title)
            title_en = self._normalize(title_en or detected_title)
            ref_standard = self._normalize(
                ref_standard or re.sub(r"^SNI\s+", "", sni_number, flags=re.I)
            )
            bsn_year = str(bsn_year or datetime.now().year)

            elements = self._build_prakata(
                doc, sni_number, title_id, title_en, ref_standard, bsn_year,
                # Engine 3 sudah menempatkan page break sebelum Introduction.
                # Karena Prakata disisipkan tepat sebelum Introduction, break
                # tersebut otomatis menjadi pemisah Daftar isi–Prakata. Untuk
                # standar tanpa Introduction, buat break di sini.
                leading_page_break=intro is None,
                # Jika Introduction tidak ada, pemisah section menuju Content
                # sudah memulai halaman baru. Page break tambahan di sini dapat
                # menghasilkan halaman kosong.
                trailing_page_break=bool(intro is not None and trailing_page_break),
            )
            for paragraph in elements:
                insertion_element.addprevious(paragraph._p)

            out_dir = os.path.dirname(os.path.abspath(output_docx))
            os.makedirs(out_dir, exist_ok=True)
            atomic_save_docx(doc, output_docx)
            check = Document(output_docx)
            if len(check.sections) != input_section_count:
                raise RuntimeError("Penyisipan Prakata mengubah jumlah section.")
            return True, output_docx
        except Exception as exc:
            return False, f"PrakataPendahuluanEngine Error: {exc}"

    def process(self, input_docx: str, output_docx: str, **kwargs):
        ok, result = self.insert_prakata(input_docx, output_docx, **kwargs)
        if not ok:
            return False, None, result
        return (
            True,
            result,
            "Prakata berhasil disisipkan setelah Daftar isi. Daftar isi, "
            "Prakata tetap berada di section 3; Introduction hanya dipertahankan jika ada pada input.",
        )

    # Alias kompatibilitas: versi baru hanya menyisipkan Prakata.
    def insert(self, input_docx: str, output_docx: str, **kwargs):
        return self.insert_prakata(input_docx, output_docx, **kwargs)
