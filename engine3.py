"""Engine 3 — sisipkan halaman Daftar isi ke keluaran Engine 2.

Urutan section hasil akhir (nomor section menggunakan hitungan Word/1-based):
1. Cover
2. Copyright
3. Daftar isi + Introduction (dipisahkan page break biasa)
4. Content sampai Bibliography

Engine ini tidak menambah section baru. Dengan demikian header, footer, margin,
dan nomor halaman milik section 3/4 dari Engine 2 tetap dipertahankan.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from typing import Optional

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


class DaftarIsiEngine:
    """Membuat atau menyisipkan satu halaman Daftar isi RSNI."""

    FONT_NAME = "Arial"
    HEADER_FONT_SIZE = 12
    FOOTER_FONT_SIZE = 10
    TITLE_FONT_SIZE = 12

    def __init__(self, font_name: str = FONT_NAME):
        self.font_name = font_name or self.FONT_NAME

    @staticmethod
    def _normalize(text: str) -> str:
        return re.sub(r"\s+", " ", text or "").strip().casefold()

    @staticmethod
    def _set_run_font(run, font_name: str, size: float, bold: bool = False):
        run.font.name = font_name
        run.font.size = Pt(size)
        run.bold = bold
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), font_name)

    @staticmethod
    def _set_page_number_format(section, fmt: str = "lowerRoman", start: int = 1):
        sect_pr = section._sectPr
        pg_num_type = sect_pr.find(qn("w:pgNumType"))
        if pg_num_type is None:
            pg_num_type = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num_type)
        pg_num_type.set(qn("w:fmt"), fmt)
        pg_num_type.set(qn("w:start"), str(start))

    @staticmethod
    def _add_page_field(paragraph, font_name: str, size: float = 10):
        run = paragraph.add_run()
        DaftarIsiEngine._set_run_font(run, font_name, size, bold=True)
        for tag, value in (
            ("w:fldChar", "begin"),
            ("w:instrText", " PAGE "),
            ("w:fldChar", "separate"),
            ("w:t", "i"),
            ("w:fldChar", "end"),
        ):
            element = OxmlElement(tag)
            if tag == "w:fldChar":
                element.set(qn("w:fldCharType"), value)
            else:
                if tag == "w:instrText":
                    element.set(qn("xml:space"), "preserve")
                element.text = value
            run._r.append(element)

    @staticmethod
    def _section_break_count_before(paragraph) -> int:
        """Hitung section break paragraf sebelum sebuah paragraf."""
        count = 0
        for sibling in paragraph._p.itersiblings(preceding=True):
            if sibling.tag != qn("w:p"):
                continue
            ppr = sibling.find(qn("w:pPr"))
            if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                count += 1
        return count

    def _new_title_paragraph(self, doc: Document):
        title_style = self._ensure_judul_style(doc)
        paragraph = doc.add_paragraph()
        paragraph.style = title_style
        paragraph.add_run("Daftar isi")
        return paragraph

    def _ensure_judul_style(self, doc: Document):
        """Pastikan style paragraf ``@Judul`` tersedia dan sesuai ketentuan."""
        try:
            style = doc.styles["@Judul"]
        except KeyError:
            style = doc.styles.add_style("@Judul", WD_STYLE_TYPE.PARAGRAPH)

        style.font.name = self.font_name
        style.font.size = Pt(self.TITLE_FONT_SIZE)
        style.font.bold = True

        # Tetapkan font untuk seluruh script Word, bukan hanya latin/ascii.
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rfonts.set(qn(f"w:{attr}"), self.font_name)

        paragraph_format = style.paragraph_format
        paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph_format.space_before = Pt(0)
        paragraph_format.space_after = Pt(0)
        paragraph_format.line_spacing = 1.0
        return style

    def _new_page_break_paragraph(self, doc: Document):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.add_run().add_break(WD_BREAK.PAGE)
        return paragraph

    def insert_after_copyright(self, input_docx: str, output_docx: str) -> str:
        """Sisipkan Daftar isi setelah Copyright, tanpa menambah section.

        Titik sisip adalah awal section 3, tepat sebelum paragraf Introduction.
        Page break biasa ditempatkan sesudah Daftar isi agar Introduction tetap
        dimulai pada halaman berikutnya tetapi masih berada di section 3.
        """
        if not input_docx or not os.path.isfile(input_docx):
            raise FileNotFoundError(f"File input tidak ditemukan: {input_docx}")
        if not output_docx:
            raise ValueError("output_docx wajib diisi")

        doc = Document(input_docx)
        input_section_count = len(doc.sections)
        if input_section_count < 4:
            raise ValueError(
                "Output Engine 2 harus memiliki minimal 4 section: Cover, "
                "Copyright, Introduction, dan Content–Bibliography. "
                f"Ditemukan {input_section_count} section."
            )

        introduction = next(
            (p for p in doc.paragraphs if self._normalize(p.text) == "introduction"),
            None,
        )
        insertion_element = introduction._p if introduction is not None else None
        if introduction is not None:
            # Dua section break sebelum Introduction berarti posisinya section 3.
            if self._section_break_count_before(introduction) != 2:
                raise ValueError("Introduction tidak berada pada awal section 3.")
        else:
            # Standar tanpa Introduction mempunyai section 3 kosong untuk
            # Daftar isi/Prakata. Sisipkan sebelum pemisah menuju Content.
            boundaries = []
            for paragraph in doc.paragraphs:
                ppr = paragraph._p.find(qn("w:pPr"))
                if ppr is not None and ppr.find(qn("w:sectPr")) is not None:
                    boundaries.append(paragraph._p)
            if len(boundaries) < 3:
                raise ValueError("Batas section front matter dan Content tidak ditemukan.")
            insertion_element = boundaries[2]

        # Idempoten: jangan menyisipkan halaman kedua jika file diproses ulang.
        if not any(self._normalize(p.text) == "daftar isi" for p in doc.paragraphs):
            title = self._new_title_paragraph(doc)
            insertion_element.addprevious(title._p)
            # Introduction berada pada section yang sama sehingga perlu page
            # break. Tanpa Introduction, sectPr menuju Content sudah nextPage.
            if introduction is not None:
                page_break = self._new_page_break_paragraph(doc)
                insertion_element.addprevious(page_break._p)

        out_dir = os.path.dirname(os.path.abspath(output_docx))
        os.makedirs(out_dir, exist_ok=True)
        doc.save(output_docx)

        # Pemeriksaan akhir: penyisipan tidak boleh mengubah jumlah section.
        checked = Document(output_docx)
        if len(checked.sections) != input_section_count:
            raise RuntimeError("Penyisipan Daftar isi mengubah jumlah section dokumen.")
        return output_docx

    def _enable_mirror_margins(self, doc: Document):
        doc.settings.odd_and_even_pages_header_footer = True
        settings = doc.settings.element
        if settings.find(qn("w:mirrorMargins")) is None:
            settings.append(OxmlElement("w:mirrorMargins"))

    def _configure_standalone_section(self, section):
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2)
        section.header_distance = Cm(1.27)
        section.footer_distance = Cm(1.27)
        self._set_page_number_format(section, "lowerRoman", 1)

    def _build_standalone_header_footer(
        self, section, sni_number: str, copyright_text: str
    ):
        section.header.is_linked_to_previous = False
        odd = section.header.paragraphs[0]
        odd.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._set_run_font(
            odd.add_run(sni_number), self.font_name, self.HEADER_FONT_SIZE, True
        )

        section.even_page_header.is_linked_to_previous = False
        even = section.even_page_header.paragraphs[0]
        even.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self._set_run_font(
            even.add_run(sni_number), self.font_name, self.HEADER_FONT_SIZE, True
        )

        section.footer.is_linked_to_previous = False
        footer = section.footer.paragraphs[0]
        width = section.page_width - section.left_margin - section.right_margin
        footer.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        self._set_run_font(
            footer.add_run(copyright_text), self.font_name, self.FOOTER_FONT_SIZE, True
        )
        footer.add_run("\t")
        self._add_page_field(footer, self.font_name, self.FOOTER_FONT_SIZE)

    def create(
        self,
        output_docx: str,
        sni_number: str = "SNI ISO xxxx-x:xxxx",
        year: Optional[int | str] = None,
        copyright_text: Optional[str] = None,
    ) -> str:
        """Tetap mendukung pembuatan Daftar isi sebagai file mandiri."""
        year = year if year is not None else datetime.now().year
        copyright_text = copyright_text or f"©BSN {year}"
        doc = Document()
        self._enable_mirror_margins(doc)
        section = doc.sections[0]
        self._configure_standalone_section(section)
        self._build_standalone_header_footer(section, sni_number, copyright_text)
        self._new_title_paragraph(doc)
        out_dir = os.path.dirname(os.path.abspath(output_docx))
        os.makedirs(out_dir, exist_ok=True)
        doc.save(output_docx)
        return output_docx

    def process(
        self,
        input_docx: Optional[str] = None,
        output_docx: Optional[str] = None,
        sni_number: str = "SNI ISO xxxx-x:xxxx",
        year: Optional[int | str] = None,
        copyright_text: Optional[str] = None,
        **kwargs,
    ):
        """Jika ada input, gabungkan; tanpa input, buat file mandiri."""
        output_docx = output_docx or "Daftar_isi_SNI.docx"
        try:
            if input_docx:
                path = self.insert_after_copyright(input_docx, output_docx)
                return (
                    True,
                    path,
                    "Daftar isi berhasil disisipkan setelah Copyright pada "
                    "section 3. Introduction hanya dipertahankan jika tersedia "
                    "pada input; Content sampai Bibliography dapat memakai "
                    "beberapa layout section.",
                )
            path = self.create(
                output_docx, sni_number, year=year, copyright_text=copyright_text
            )
            return True, path, "Daftar isi berhasil dibuat."
        except Exception as exc:
            return False, None, f"{type(exc).__name__}: {exc}"