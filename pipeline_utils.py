"""Shared reliability helpers for Generator RSNI engines.

No external service/database is required.  The helpers deliberately stay small:
validate OOXML packages, save python-docx documents atomically, and expose a
cheap structural preflight that every engine can reuse.
"""
from __future__ import annotations
import os, tempfile, zipfile
from pathlib import Path
from docx import Document
from lxml import etree

_REQUIRED = ('[Content_Types].xml', '_rels/.rels', 'word/document.xml')

def validate_docx(path: str, *, require_paragraphs: bool = True) -> dict:
    if not path or not os.path.isfile(path):
        raise FileNotFoundError(f'DOCX tidak ditemukan: {path}')
    if not zipfile.is_zipfile(path):
        raise ValueError(f'File bukan paket DOCX/ZIP yang valid: {path}')
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        missing = [n for n in _REQUIRED if n not in names]
        if missing:
            raise ValueError('Komponen DOCX wajib hilang: ' + ', '.join(missing))
        bad = z.testzip()
        if bad:
            raise ValueError(f'CRC DOCX rusak pada: {bad}')
        etree.fromstring(z.read('word/document.xml'))
    doc = Document(path)
    if not doc.sections:
        raise ValueError('DOCX tidak memiliki section Word yang valid.')
    if require_paragraphs and not doc.paragraphs and not doc.tables:
        raise ValueError('DOCX tidak memiliki konten body yang dapat diproses.')
    return {'sections': len(doc.sections), 'paragraphs': len(doc.paragraphs), 'tables': len(doc.tables)}

def atomic_save_docx(doc: Document, output_path: str) -> str:
    """Save then validate, replacing the destination only after success."""
    output_path = os.path.abspath(output_path)
    parent = os.path.dirname(output_path) or '.'
    os.makedirs(parent, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix='.rsni-', suffix='.docx', dir=parent)
    os.close(fd)
    try:
        doc.save(tmp)
        validate_docx(tmp)
        os.replace(tmp, output_path)
        return output_path
    finally:
        try:
            if os.path.exists(tmp): os.remove(tmp)
        except OSError:
            pass

def atomic_replace_file(tmp_path: str, output_path: str) -> str:
    validate_docx(tmp_path)
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    os.replace(tmp_path, output_path)
    return output_path
