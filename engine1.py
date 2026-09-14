"""Engine 1 - trim dokumen ISO menjadi Introduction + Content.

Dokumen dipotong langsung pada XML Word agar style, numbering, gambar,
header/footer, section, dan pengaturan halaman asli tetap dipertahankan.
"""

from __future__ import annotations

import os
import re
from typing import Optional, Tuple

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


class IntroductionContentTrimmerEngine:
    INTRO_RE = re.compile(r"^\s*introduction\s*$", re.IGNORECASE)
    # Edisi IEC lama memakai heading resmi "Scope and object". Keduanya
    # merupakan Pasal 1 dan harus diterima sebagai awal Content.
    SCOPE_RE = re.compile(
        r"^\s*(?:\d+(?:\.0)?\s+)?scope(?:\s+and\s+object)?\s*$",
        re.IGNORECASE,
    )
    BIBLIO_RE = re.compile(r"^\s*bibliograph(?:y|ie)\s*$", re.IGNORECASE)
    NUMBER_PREFIX_RE = re.compile(
        r"^\s*(?:Annex\s+[A-Z]|[A-Z](?:\.\d+)+|\d+(?:\.\d+)*)\s+",
        re.IGNORECASE,
    )

    @staticmethod
    def _text(paragraph) -> str:
        return " ".join(paragraph.text.replace("\xa0", " ").split())

    @staticmethod
    def _has_sect_pr(paragraph) -> bool:
        return bool(paragraph._p.xpath("./w:pPr/w:sectPr"))

    @staticmethod
    def _remove_element(element) -> None:
        parent = element.getparent()
        if parent is not None:
            parent.remove(element)

    @staticmethod
    def _format_unlinked_run(run_element) -> None:
        """Format teks bekas hyperlink: Arial 11, hitam, tanpa underline."""
        r_pr = run_element.find(qn("w:rPr"))
        if r_pr is None:
            r_pr = OxmlElement("w:rPr")
            run_element.insert(0, r_pr)

        # Hapus style karakter Hyperlink dan seluruh properti yang dapat
        # mewariskan warna biru, underline, font, atau ukuran lama.
        for tag in ("w:rStyle", "w:rFonts", "w:color", "w:u", "w:sz", "w:szCs"):
            for old in list(r_pr.findall(qn(tag))):
                r_pr.remove(old)

        fonts = OxmlElement("w:rFonts")
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{attr}"), "Arial")
        r_pr.append(fonts)

        color = OxmlElement("w:color")
        color.set(qn("w:val"), "000000")
        color.set(qn("w:themeColor"), "text1")
        r_pr.append(color)

        underline = OxmlElement("w:u")
        underline.set(qn("w:val"), "none")
        r_pr.append(underline)

        for tag in ("w:sz", "w:szCs"):
            size = OxmlElement(tag)
            size.set(qn("w:val"), "22")  # half-points: 11 pt
            r_pr.append(size)

    @classmethod
    def _unlink_container(cls, container) -> int:
        """Ubah seluruh jenis hyperlink dalam satu part Word menjadi teks biasa."""
        removed = 0

        # Hyperlink normal: pindahkan isi run ke posisi pembungkus link.
        for hyperlink in list(container.xpath(".//w:hyperlink")):
            parent = hyperlink.getparent()
            if parent is None:
                continue
            index = parent.index(hyperlink)
            children = list(hyperlink)
            for child in children:
                runs = [child] if child.tag == qn("w:r") else list(child.iter(qn("w:r")))
                for run in runs:
                    cls._format_unlinked_run(run)
                parent.insert(index, child)
                index += 1
            parent.remove(hyperlink)
            removed += 1

        # Simple field: <w:fldSimple w:instr="HYPERLINK ...">.
        for field in list(container.xpath(".//w:fldSimple")):
            instruction = field.get(qn("w:instr")) or ""
            if "HYPERLINK" not in instruction.upper():
                continue
            parent = field.getparent()
            if parent is None:
                continue
            index = parent.index(field)
            for child in list(field):
                runs = [child] if child.tag == qn("w:r") else list(child.iter(qn("w:r")))
                for run in runs:
                    cls._format_unlinked_run(run)
                parent.insert(index, child)
                index += 1
            parent.remove(field)
            removed += 1

        # Complex field: begin → instrText HYPERLINK → separate → result → end.
        for paragraph in container.xpath(".//w:p"):
            children = list(paragraph)
            cursor = 0
            while cursor < len(children):
                begin = children[cursor].find(qn("w:fldChar"))
                if begin is None or begin.get(qn("w:fldCharType")) != "begin":
                    cursor += 1
                    continue

                end_index = None
                separate_index = None
                instruction_parts = []
                depth = 0
                for index in range(cursor, len(children)):
                    child = children[index]
                    field_char = child.find(qn("w:fldChar"))
                    if field_char is not None:
                        kind = field_char.get(qn("w:fldCharType"))
                        if kind == "begin":
                            depth += 1
                        elif kind == "separate" and depth == 1:
                            separate_index = index
                        elif kind == "end":
                            depth -= 1
                            if depth == 0:
                                end_index = index
                                break
                    instruction_parts.extend(
                        node.text or "" for node in child.findall(qn("w:instrText"))
                    )

                instruction = "".join(instruction_parts).upper()
                if (
                    end_index is None
                    or separate_index is None
                    or "HYPERLINK" not in instruction
                ):
                    cursor += 1
                    continue

                for result in children[separate_index + 1:end_index]:
                    runs = [result] if result.tag == qn("w:r") else list(result.iter(qn("w:r")))
                    for run in runs:
                        cls._format_unlinked_run(run)
                for index in list(range(cursor, separate_index + 1)) + [end_index]:
                    child = children[index]
                    if child.getparent() is paragraph:
                        paragraph.remove(child)
                removed += 1
                children = list(paragraph)
                cursor = 0

        return removed

    @classmethod
    def _remove_all_hyperlinks(cls, doc: Document) -> int:
        """Hapus hyperlink di body, header, dan footer tanpa menghapus teksnya."""
        removed = cls._unlink_container(doc.element.body)
        seen_parts = set()
        for section in doc.sections:
            for container in (
                section.header,
                section.first_page_header,
                section.even_page_header,
                section.footer,
                section.first_page_footer,
                section.even_page_footer,
            ):
                part_key = str(container.part.partname)
                if part_key in seen_parts:
                    continue
                seen_parts.add(part_key)
                removed += cls._unlink_container(container._element)
        return removed

    @staticmethod
    def _set_page_start(section, number: int = 1) -> None:
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num)
        pg_num.set(qn("w:start"), str(number))

    @staticmethod
    def _set_style_font(style, name: str, size: int, bold: bool) -> None:
        style.font.name = name
        style.font.size = Pt(size)
        style.font.bold = bold
        r_pr = style.element.get_or_add_rPr()
        r_fonts = r_pr.get_or_add_rFonts()
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            r_fonts.set(qn(f"w:{attr}"), name)

    def _ensure_output_styles(self, doc: Document):
        """Buat/perbarui dua style khusus tanpa bergantung template input."""
        styles = doc.styles
        try:
            judul = styles["@Judul"]
        except KeyError:
            judul = styles.add_style("@Judul", WD_STYLE_TYPE.PARAGRAPH)
        self._set_style_font(judul, "Arial", 12, True)
        judul.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        judul.paragraph_format.line_spacing = 1.0
        judul.paragraph_format.space_before = Pt(0)
        judul.paragraph_format.space_after = Pt(0)

        try:
            pasal = styles["@Pasal"]
        except KeyError:
            pasal = styles.add_style("@Pasal", WD_STYLE_TYPE.PARAGRAPH)
        self._set_style_font(pasal, "Arial", 11, True)
        pasal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pasal.paragraph_format.line_spacing = 1.0
        pasal.paragraph_format.space_before = Pt(0)
        pasal.paragraph_format.space_after = Pt(0)
        return judul, pasal

    @staticmethod
    def _replace_paragraph_text(paragraph, text: str) -> None:
        """Ganti teks tetapi pertahankan pPr/section/page-break paragraf."""
        paragraph.clear()
        paragraph.add_run(text)

    @staticmethod
    def _section_starts(doc: Document):
        """Kembalikan pasangan (section, indeks paragraf awal section)."""
        starts = [0]
        for idx, paragraph in enumerate(doc.paragraphs):
            if IntroductionContentTrimmerEngine._has_sect_pr(paragraph):
                starts.append(idx + 1)
        # Dokumen Word yang valid mempunyai satu sectPr per section. Batasi
        # hasil untuk berjaga-jaga terhadap file hasil konversi yang aneh.
        return list(zip(doc.sections, starts[:len(doc.sections)]))

    def _numbered_content_start(self, doc: Document) -> Optional[int]:
        """Cari awal section halaman Arab yang secara eksplisit mulai dari 1.

        Dokumen amendment/corrigendum sering tidak mempunyai Pasal 1 Scope.
        ISO menandai badan dokumen tersebut dengan section bernomor halaman 1;
        front matter sebelumnya tidak memiliki ``w:start=1`` atau memakai
        angka Romawi. Ini menjadi fallback yang lebih aman daripada menebak
        isi berdasarkan kata tertentu.
        """
        candidates = []
        for section, start_idx in self._section_starts(doc):
            pg_num = section._sectPr.find(qn("w:pgNumType"))
            if pg_num is None or pg_num.get(qn("w:start")) != "1":
                continue
            if (pg_num.get(qn("w:fmt")) or "decimal").casefold() != "decimal":
                continue
            candidates.append(start_idx)
        return candidates[-1] if candidates else None

    @classmethod
    def _fallback_cover_title(cls, doc: Document, content_start: int) -> str:
        """Ambil judul Inggris dari cover jika field judul Content rusak."""
        candidates = []
        for index, paragraph in enumerate(doc.paragraphs[:content_start]):
            text = cls._text(paragraph)
            if not text or 'reference source not found' in text.casefold():
                continue
            style = paragraph.style.name.casefold() if paragraph.style else ''
            if 'cover' not in style:
                continue
            if len(text) >= 20 and ('—' in text or '-' in text):
                candidates.append((index, text))
        # Cover ISO menempatkan judul Inggris sebelum judul Prancis.
        return candidates[0][1] if candidates else ''

    def _format_trimmed_document(
        self, doc: Document, fallback_title: str = ''
    ) -> None:
        judul, pasal = self._ensure_output_styles(doc)

        # Cari Scope pada dokumen yang sudah dipotong, lalu anggap semua teks
        # nonkosong sesudah section break terakhir dan sebelum Scope sebagai
        # judul Content. Beberapa dokumen memecah judul menjadi 2-3 paragraf.
        _, scope_idx, _ = self._find_markers(doc)
        amendment_mode = scope_idx is None
        content_start = self._content_start(doc, scope_idx)
        if scope_idx is not None:
            title_paragraphs = [
                p for p in doc.paragraphs[content_start:scope_idx] if self._text(p)
            ]
        else:
            # Pada amendment, paragraf pertama halaman 1 adalah judul standar.
            # Jangan gabungkan instruksi perubahan sesudahnya ke dalam judul.
            title_paragraphs = []
            for paragraph in doc.paragraphs[content_start:]:
                if self._text(paragraph):
                    title_paragraphs = [paragraph]
                    break
        if title_paragraphs:
            title_text = " ".join(self._text(p) for p in title_paragraphs)
            title_text = re.sub(r"\s+", " ", title_text).strip()
            if (
                'reference source not found' in title_text.casefold()
                and fallback_title
            ):
                title_text = fallback_title
            first_title = title_paragraphs[0]
            self._replace_paragraph_text(first_title, title_text)
            first_title.style = judul
            for extra_title in title_paragraphs[1:]:
                self._remove_element(extra_title._p)

        main_counters = [0] * 10
        annex_counters = [0] * 10
        annex_index = 0
        annex_letter = ""
        in_annex = False

        # List dibuat ulang karena penggabungan judul di atas mengubah koleksi.
        for paragraph in list(doc.paragraphs):
            text = self._text(paragraph)
            if not text:
                continue
            style_name = paragraph.style.name
            lower_text = text.lower()

            if self.INTRO_RE.fullmatch(text) or self.BIBLIO_RE.fullmatch(text):
                self._replace_paragraph_text(paragraph, text)
                paragraph.style = judul
                if self.BIBLIO_RE.fullmatch(text):
                    # Bibliography dimulai pada halaman baru, tetapi tetap
                    # berada dalam section yang sama agar header/footer dan
                    # kelanjutan nomor halaman tidak berubah.
                    paragraph.paragraph_format.page_break_before = True
                continue

            # Amendment/corrigendum berisi instruksi seperti "Clause 5" dan
            # "Annex B" yang menunjuk bagian dokumen induk, bukan struktur
            # heading baru. Pertahankan teks/style/nomornya apa adanya.
            if amendment_mode:
                continue

            # Jangan menganggap kalimat naratif seperti ``Annex A lists ...``
            # sebagai heading Annex. ISO/IEC 9797-2 memuat tiga kalimat seperti
            # itu di Pasal 5; deteksi lama membuatnya menjadi Annex A/B/C lalu
            # seluruh pasal sesudahnya salah diberi nomor C.x dan Annex asli
            # bergeser/hilang. Heading sah ditandai style ANNEX atau teks polos
            # yang tepat berupa ``Annex A`` / ``Annex A (normative|informative)``.
            plain_annex_heading = re.match(
                r"^annex\s+[a-z](?:\s*$|\s+\((?:normative|informative)\)"
                r"(?:\s+.*)?$)",
                text,
                re.I,
            )
            if style_name.casefold() == "annex" or plain_annex_heading:
                in_annex = True
                annex_index += 1
                annex_letter = chr(ord("A") + annex_index - 1)
                annex_counters = [0] * 10
                cleaned = re.sub(r"^Annex\s+[A-Z]\s*", "", text, flags=re.I)
                informative = ""
                match = re.match(r"^\s*(\([^)]*\))\s*(.*)$", cleaned)
                if match:
                    informative, cleaned = match.group(1), match.group(2)
                parts = [f"Annex {annex_letter}"]
                if informative:
                    parts.append(informative)
                if cleaned.strip():
                    parts.append(cleaned.strip())
                self._replace_paragraph_text(paragraph, "\n".join(parts))
                paragraph.style = judul
                # Annex selalu dimulai pada halaman baru tanpa membuat
                # section baru. Properti ini idempoten dan tidak menambahkan
                # paragraf/page break ganda saat engine dijalankan ulang.
                paragraph.paragraph_format.page_break_before = True
                continue

            heading_match = re.fullmatch(r"Heading\s+(\d+)", style_name, re.I)
            annex_heading_match = re.fullmatch(r"a(\d+)", style_name, re.I)
            if not heading_match and not annex_heading_match:
                continue

            level = int((annex_heading_match or heading_match).group(1))
            cleaned = self.NUMBER_PREFIX_RE.sub("", text).strip()

            if in_annex or annex_heading_match:
                in_annex = True
                if not annex_letter:
                    annex_index = 1
                    annex_letter = "A"
                logical_level = max(1, level - 1)
                annex_counters[logical_level] += 1
                for idx in range(logical_level + 1, len(annex_counters)):
                    annex_counters[idx] = 0
                suffix = ".".join(
                    str(annex_counters[idx])
                    for idx in range(1, logical_level + 1)
                )
                number = f"{annex_letter}.{suffix}"
            else:
                level = max(1, min(level, 9))
                main_counters[level] += 1
                for idx in range(level + 1, len(main_counters)):
                    main_counters[idx] = 0
                number = ".".join(
                    str(main_counters[idx]) for idx in range(1, level + 1)
                )

            # Empat spasi literal memastikan nomor tetap terlihat sekalipun
            # style numbering Word diganti menjadi @Pasal.
            self._replace_paragraph_text(paragraph, f"{number}    {cleaned}")
            paragraph.style = pasal

    def _find_markers(
        self, doc: Document
    ) -> Tuple[Optional[int], Optional[int], Optional[int]]:
        intro_idx = None
        scope_idx = None
        biblio_idx = None

        for idx, paragraph in enumerate(doc.paragraphs):
            text = self._text(paragraph)
            if intro_idx is None and self.INTRO_RE.fullmatch(text):
                intro_idx = idx
            if scope_idx is None and self.SCOPE_RE.fullmatch(text):
                scope_idx = idx
            if biblio_idx is None and self.BIBLIO_RE.fullmatch(text):
                biblio_idx = idx

        if scope_idx is None and self._numbered_content_start(doc) is None:
            raise ValueError(
                "Halaman Content tidak ditemukan: pasal 'Scope' tidak ada dan "
                "section dengan penomoran halaman mulai dari 1 tidak ditemukan."
            )
        if biblio_idx is not None and scope_idx is not None and biblio_idx < scope_idx:
            biblio_idx = None
        return intro_idx, scope_idx, biblio_idx

    def _content_start(self, doc: Document, scope_idx: Optional[int]) -> int:
        """Ambil awal section Content, termasuk judul standar sebelum Scope."""
        if scope_idx is None:
            numbered_start = self._numbered_content_start(doc)
            if numbered_start is None:
                raise ValueError("Awal halaman Content bernomor 1 tidak ditemukan.")
            start = numbered_start
            while start < len(doc.paragraphs) and not self._text(doc.paragraphs[start]):
                start += 1
            return start

        paragraphs = doc.paragraphs
        boundary = -1
        for idx in range(scope_idx - 1, -1, -1):
            if self._has_sect_pr(paragraphs[idx]):
                boundary = idx
                break

        start = boundary + 1
        while start < scope_idx and not self._text(paragraphs[start]):
            start += 1
        return start

    def process(self, input_docx: str, output_docx: str) -> Tuple[bool, str]:
        try:
            if not os.path.isfile(input_docx):
                raise FileNotFoundError(f"File tidak ditemukan: {input_docx}")

            doc = Document(input_docx)
            intro_idx, scope_idx, biblio_idx = self._find_markers(doc)
            content_start = self._content_start(doc, scope_idx)
            fallback_title = self._fallback_cover_title(doc, content_start)

            # Introduction hanya sah bila berada sebelum Content. Jika tidak,
            # dokumen dianggap tidak memiliki halaman Introduction terpisah.
            has_intro = intro_idx is not None and intro_idx < content_start
            start_idx = intro_idx if has_intro else content_start

            paragraphs = doc.paragraphs
            for paragraph in paragraphs[:start_idx]:
                self._remove_element(paragraph._p)

            # Buang tabel/objek tingkat-body yang berada sebelum paragraf awal.
            # Umumnya dokumen ISO memakai paragraf; loop ini juga menangani
            # cover yang dibangun sebagai tabel tanpa mengubah isi sesudahnya.
            body = doc._element.body
            first_kept = doc.paragraphs[0]._p if doc.paragraphs else None
            if first_kept is not None:
                for child in list(body):
                    if child is first_kept:
                        break
                    if child.tag != qn("w:sectPr"):
                        self._remove_element(child)

            if not has_intro:
                # Content menjadi halaman pertama dan harus dimulai dari 1.
                self._set_page_start(doc.sections[0], 1)

            # Hyperlink harus dibongkar sebelum formatting paragraf. Run di
            # dalam w:hyperlink tidak selalu muncul pada paragraph.runs.
            hyperlinks_removed = self._remove_all_hyperlinks(doc)
            self._format_trimmed_document(doc, fallback_title=fallback_title)

            os.makedirs(os.path.dirname(os.path.abspath(output_docx)), exist_ok=True)
            doc.save(output_docx)

            intro_msg = "Introduction + Content" if has_intro else "Content (mulai halaman 1)"
            biblio_msg = " sampai Bibliography" if biblio_idx is not None else ""
            link_msg = f" {hyperlinks_removed} hyperlink diubah menjadi teks biasa."
            return True, f"Trim berhasil: {intro_msg}{biblio_msg}.{link_msg}"
        except Exception as exc:
            return False, str(exc)


# Alias ringkas untuk kompatibilitas pemanggilan dari app.
Engine1 = IntroductionContentTrimmerEngine
